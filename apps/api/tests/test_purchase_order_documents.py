"""PO documents (migration 0059): one PO groups open orders of one supplier.

The PO number is its own running number (`po_document_no_seq`); orders on it are marked ordered;
the delivery charge is a real Delivery Fee order; totals skip Cancelled / Rejected lines.
"""
import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.db import SessionLocal

from .conftest import truncate_fixture
from .helpers import login, login_same_workspace

_cleanup = truncate_fixture("po_document", "vendor_contact", "po_invoice")


def _sql(sql: str, **p):
    s = SessionLocal()
    try:
        out = s.execute(text(sql), p)
        val = out.scalar() if out.returns_rows else None
        s.commit()
        return val
    finally:
        s.close()


@pytest.fixture
def ws():
    buyer, wid, _ = login("purchase_officer", prefix="pd")
    v1 = _sql("INSERT INTO vendors(name, category, workspace_id) VALUES ('Laminex', 'Board', :w)"
              " RETURNING vendor_id", w=wid)
    v2 = _sql("INSERT INTO vendors(name, category, workspace_id) VALUES ('Other', 'Board', :w)"
              " RETURNING vendor_id", w=wid)
    viewer, _ = login_same_workspace(wid, "viewer")
    return {"wid": wid, "buyer": buyer, "viewer": viewer, "v1": v1, "v2": v2}


def _order(ws, vendor="v1", total="100.00", **extra):
    body = {"vendor_id": ws[vendor], "description": "Board", "quantity": 1, "unit_cost": total,
            "total_amount": total, **extra}
    r = ws["buyer"].post("/orders", json=body)
    assert r.status_code == 201, r.text
    return r.json()["po_id"]


def _po(ws, *order_ids, **body):
    r = ws["buyer"].post("/purchase-orders", json={"order_ids": list(order_ids), **body})
    assert r.status_code == 201, r.text
    return r.json()


def _code(r):
    return r.json()["detail"]["code"]


def test_create_marks_orders_ordered_and_totals(ws):
    a, b = _order(ws, total="100.00"), _order(ws, total="50.00")
    po = _po(ws, a, b, reference_no="R-1")
    assert po["po_no"] >= 193971  # 0060: after the old system's last PO, 193970
    assert po["status"] == "Open" and po["reference_no"] == "R-1"
    assert len(po["lines"]) == 2
    assert po["totals"]["total_ex"] == "150.00"
    assert po["totals"]["total_inc"] == "165.00"
    o = ws["buyer"].get(f"/orders/{a}").json()
    assert o["po_no"] == po["po_no"] and o["date_ordered"] is not None


def test_rules(ws):
    a, other = _order(ws), _order(ws, vendor="v2")
    assert _code(ws["buyer"].post("/purchase-orders", json={"order_ids": [a, other]})) == "SUPPLIER_MISMATCH"
    _po(ws, a)
    assert _code(ws["buyer"].post("/purchase-orders", json={"order_ids": [a]})) == "ORDER_ALREADY_ON_PO"
    done = _order(ws)
    assert ws["buyer"].patch(f"/orders/{done}", json={"status": "Cancelled"}).status_code == 200
    assert _code(ws["buyer"].post("/purchase-orders", json={"order_ids": [done]})) == "ORDER_NOT_OPEN"
    assert ws["buyer"].post("/purchase-orders", json={"order_ids": [10**9]}).status_code == 404
    assert ws["buyer"].post("/purchase-orders", json={"order_ids": []}).status_code == 422


def test_add_remove_and_vendor_guard(ws):
    a, b = _order(ws), _order(ws)
    po = _po(ws, a)
    id_ = po["po_document_id"]
    r = ws["buyer"].post(f"/purchase-orders/{id_}/orders", json={"order_ids": [b]})
    assert r.status_code == 200 and len(r.json()["lines"]) == 2
    r = ws["buyer"].patch(f"/orders/{a}", json={"vendor_id": ws["v2"]})
    assert r.status_code == 409 and _code(r) == "ORDER_ON_PO"
    r = ws["buyer"].delete(f"/purchase-orders/{id_}/orders/{b}")
    assert r.status_code == 200 and len(r.json()["lines"]) == 1
    assert ws["buyer"].get(f"/orders/{b}").json()["po_no"] is None
    assert ws["buyer"].delete(f"/purchase-orders/{id_}/orders/{b}").status_code == 404


def test_delivery_is_a_real_order(ws):
    a = _order(ws)
    id_ = _po(ws, a)["po_document_id"]
    r = ws["buyer"].put(f"/purchase-orders/{id_}/delivery", json={"amount": "45.00"})
    assert r.status_code == 200, r.text
    po = r.json()
    assert po["delivery_charge"] == "45.00" and po["totals"]["total_ex"] == "145.00"
    assert po["lines"][-1]["is_delivery"] is True
    d = po["delivery_order_id"]
    assert _code(ws["buyer"].delete(f"/purchase-orders/{id_}/orders/{d}")) == "IS_DELIVERY_ORDER"
    po = ws["buyer"].put(f"/purchase-orders/{id_}/delivery", json={"amount": "60.00"}).json()
    assert po["delivery_order_id"] == d and po["delivery_charge"] == "60.00"
    po = ws["buyer"].delete(f"/purchase-orders/{id_}/delivery").json()
    assert po["delivery_order_id"] is None and po["totals"]["total_ex"] == "100.00"
    assert _code(ws["buyer"].delete(f"/purchase-orders/{id_}/delivery")) == "NO_DELIVERY"


def test_totals_skip_cancelled_lines(ws):
    a, b = _order(ws, total="100.00"), _order(ws, total="50.00")
    id_ = _po(ws, a, b)["po_document_id"]
    assert ws["buyer"].patch(f"/orders/{b}", json={"status": "Cancelled"}).status_code == 200
    assert ws["buyer"].get(f"/purchase-orders/{id_}").json()["totals"]["total_ex"] == "100.00"


def test_cancel_frees_orders(ws):
    a = _order(ws)
    id_ = _po(ws, a)["po_document_id"]
    ws["buyer"].put(f"/purchase-orders/{id_}/delivery", json={"amount": "10"})
    r = ws["buyer"].post(f"/purchase-orders/{id_}/cancel")
    assert r.status_code == 200 and r.json()["status"] == "Cancelled"
    assert ws["buyer"].get(f"/orders/{a}").json()["po_no"] is None
    assert _code(ws["buyer"].post(f"/purchase-orders/{id_}/orders", json={"order_ids": [a]})) == "PO_CANCELLED"
    _po(ws, a)  # free again


def test_patch_header(ws):
    id_ = _po(ws, _order(ws))["po_document_id"]
    r = ws["buyer"].patch(f"/purchase-orders/{id_}", json={"memo": "hi", "deliver_to": "Site 1"})
    assert r.status_code == 200 and r.json()["memo"] == "hi"
    assert ws["buyer"].patch(f"/purchase-orders/{id_}", json={"date_ordered": None}).status_code == 422


def test_invoices(ws):
    a = _order(ws)
    id_ = _po(ws, a)["po_document_id"]
    r = ws["buyer"].post(f"/purchase-orders/{id_}/invoices",
                         json={"order_id": a, "invoice_number": "INV-1", "total_inc": "110.00"})
    assert r.status_code == 201, r.text
    po = r.json()
    assert po["lines"][0]["invoiced"] is True and po["invoices"][0]["invoice_number"] == "INV-1"
    inv = po["invoices"][0]["invoice_id"]
    r = ws["buyer"].patch(f"/purchase-orders/{id_}/invoices/{inv}", json={"invoice_number": "INV-2"})
    assert r.json()["invoices"][0]["invoice_number"] == "INV-2"
    assert ws["buyer"].patch(f"/purchase-orders/{id_}/invoices/{inv}",
                             json={"invoice_number": None}).status_code == 422
    assert _code(ws["buyer"].delete(f"/purchase-orders/{id_}/orders/{a}")) == "ORDER_HAS_INVOICES"
    r = ws["buyer"].delete(f"/purchase-orders/{id_}/invoices/{inv}")
    assert r.status_code == 200 and r.json()["invoices"] == []
    assert ws["buyer"].post(f"/purchase-orders/{id_}/invoices",
                            json={"order_id": _order(ws), "invoice_number": "X"}).status_code == 404


def test_outputs(ws):
    id_ = _po(ws, _order(ws), reference_no="R9")["po_document_id"]
    r = ws["buyer"].get(f"/purchase-orders/{id_}/pdf")
    assert r.status_code == 200 and r.content[:4] == b"%PDF"
    r = ws["buyer"].get(f"/purchase-orders/{id_}/authorisation-form")
    assert r.status_code == 200 and r.content[:4] == b"%PDF"
    r = ws["buyer"].get(f"/purchase-orders/{id_}/excel")
    assert r.status_code == 200
    cells = [str(c.value) for row in load_workbook(io.BytesIO(r.content)).active.iter_rows() for c in row]
    assert any("R9" in c for c in cells)


def test_list(ws):
    po = _po(ws, _order(ws))
    rows = ws["buyer"].get("/purchase-orders").json()["purchase_orders"]
    assert [r["po_no"] for r in rows] == [po["po_no"]] and rows[0]["line_count"] == 1


def test_rbac(ws):
    a = _order(ws)
    assert ws["viewer"].post("/purchase-orders", json={"order_ids": [a]}).status_code == 403
    id_ = _po(ws, a)["po_document_id"]
    assert ws["viewer"].get(f"/purchase-orders/{id_}").status_code in (200, 403)
    assert ws["viewer"].post(f"/purchase-orders/{id_}/cancel").status_code == 403


def test_other_workspace_is_404(ws):
    """Every PO operation answers another workspace's admin like a nonexistent id."""
    a = _order(ws)
    id_ = _po(ws, a)["po_document_id"]
    inv = ws["buyer"].post(f"/purchase-orders/{id_}/invoices",
                           json={"order_id": a, "invoice_number": "I"}).json()["invoices"][0]["invoice_id"]
    other, _, _ = login("admin", prefix="pe")
    ghost = 10**9
    calls = [
        ("get", "", None), ("patch", "", {"memo": "x"}), ("post", "/cancel", None),
        ("get", "/pdf", None), ("get", "/authorisation-form", None), ("get", "/excel", None),
        ("post", "/orders", {"order_ids": [a]}), ("delete", f"/orders/{a}", None),
        ("put", "/delivery", {"amount": "1"}), ("delete", "/delivery", None),
        ("post", "/invoices", {"order_id": a, "invoice_number": "Z"}),
        ("patch", f"/invoices/{inv}", {"invoice_number": "Z"}), ("delete", f"/invoices/{inv}", None),
    ]
    for method, tail, body in calls:
        real = getattr(other, method)(f"/purchase-orders/{id_}{tail}", **({"json": body} if body else {}))
        gone = getattr(other, method)(f"/purchase-orders/{ghost}{tail.replace(str(a), str(ghost)).replace(str(inv), str(ghost))}",
                                      **({"json": body} if body else {}))
        assert real.status_code == gone.status_code == 404, (method, tail, real.status_code, gone.status_code)
    # foreign ids in a body
    mine = _order(ws)
    po2 = _po(ws, mine)["po_document_id"]
    assert other.post("/purchase-orders", json={"order_ids": [a]}).status_code == 404
    assert ws["buyer"].post("/purchase-orders", json={"order_ids": [_order(ws)], "ordered_by_id": _other_user(other)}).status_code == 404
    assert ws["buyer"].patch(f"/purchase-orders/{po2}", json={"ordered_by_id": _other_user(other)}).status_code == 404
    assert other.get(f"/suppliers/{ws['v1']}/contacts").status_code == 404
    assert other.get(f"/purchase-orders?vendor_id={ws['v1']}").json() == {"purchase_orders": []}
    assert ws["buyer"].get(f"/purchase-orders/{id_}").status_code == 200  # nothing changed


def _other_user(client) -> int:
    return client.get("/auth/me").json()["id"]


def test_supplier_contacts(ws):
    c, v = ws["buyer"], ws["v1"]
    r = c.post(f"/suppliers/{v}/contacts", json={"name": " Sam ", "email": "sam@x.test", "position": "Rep"})
    assert r.status_code == 201 and r.json()["name"] == "Sam"
    cid = r.json()["contact_id"]
    assert c.post(f"/suppliers/{v}/contacts", json={"name": "Bad", "email": "nope"}).status_code == 422
    assert c.post(f"/suppliers/{v}/contacts", json={"name": " "}).status_code == 422
    assert c.patch(f"/suppliers/{v}/contacts/{cid}", json={"name": None}).status_code == 422
    r = c.patch(f"/suppliers/{v}/contacts/{cid}", json={"mobile": "0400"})
    assert r.json()["mobile"] == "0400"
    assert ws["viewer"].post(f"/suppliers/{v}/contacts", json={"name": "N"}).status_code == 403
    assert [x["contact_id"] for x in c.get(f"/suppliers/{v}/contacts").json()["contacts"]] == [cid]
    assert c.delete(f"/suppliers/{v}/contacts/{cid}").status_code == 204
    assert c.delete(f"/suppliers/{v}/contacts/{cid}").status_code == 404


def test_staff_emails_come_from_the_orderbook_route_only(ws):
    """The PO panel's Cc list: names + emails for Orderbook readers; /workspace/team stays email-free."""
    staff = ws["buyer"].get("/purchase-orders/staff").json()["staff"]
    assert staff and all(x["email"] and x["full_name"] for x in staff)
    assert all("email" not in m for m in ws["buyer"].get("/workspace/team").json()["members"])
    other, _, _ = login("admin", prefix="ps")
    mine = {x["id"] for x in staff}
    assert mine.isdisjoint({x["id"] for x in other.get("/purchase-orders/staff").json()["staff"]})


# ── code-review fixes (October 2026) ───────────────────────────────────────────────────────────

def _order_row(c, order):
    return c.get(f"/orders/{order}").json()


def test_a_hand_typed_date_ordered_survives_leaving_a_po_with_the_same_date(ws):
    """The PO only clears a Date ordered it set itself, not one that happens to equal its own."""
    c = ws["buyer"]
    a = _order(ws)
    assert c.patch(f"/orders/{a}", json={"date_ordered": "2020-05-05"}).status_code == 200
    assert _order_row(c, a)["date_ordered"] == "2020-05-05"
    id_ = _po(ws, a)["po_document_id"]
    assert c.patch(f"/purchase-orders/{id_}", json={"date_ordered": "2020-05-05"}).status_code == 200
    c.delete(f"/purchase-orders/{id_}/orders/{a}")
    assert _order_row(c, a)["date_ordered"] == "2020-05-05"


def test_the_po_date_is_cleared_even_after_the_po_date_was_edited(ws):
    c = ws["buyer"]
    a = _order(ws)                                             # no date: the PO sets it
    id_ = _po(ws, a)["po_document_id"]
    assert _order_row(c, a)["date_ordered"] is not None
    assert c.patch(f"/purchase-orders/{id_}", json={"date_ordered": "2020-01-01"}).status_code == 200
    c.delete(f"/purchase-orders/{id_}/orders/{a}")
    assert _order_row(c, a)["date_ordered"] is None            # back in "to be ordered"


def test_editing_an_orders_date_by_hand_keeps_it_when_it_leaves_the_po(ws):
    c = ws["buyer"]
    a = _order(ws)
    id_ = _po(ws, a)["po_document_id"]
    assert c.patch(f"/orders/{a}", json={"date_ordered": "2021-02-03"}).status_code == 200
    c.delete(f"/purchase-orders/{id_}/orders/{a}")
    assert _order_row(c, a)["date_ordered"] == "2021-02-03"


def test_cancelling_a_po_leaves_a_delivered_orders_date_and_versions_the_rest(ws):
    c = ws["buyer"]
    a, b = _order(ws), _order(ws)
    id_ = _po(ws, a, b)["po_document_id"]
    assert c.patch(f"/orders/{a}", json={"status": "Delivered"}).status_code == 200
    before = _order_row(c, a)["date_ordered"]
    vb = _order_row(c, b)["field_versions"].get("date_ordered", 0)
    assert c.post(f"/purchase-orders/{id_}/cancel").status_code == 200
    assert _order_row(c, a)["date_ordered"] == before          # delivered: history stays
    after_b = _order_row(c, b)
    assert after_b["date_ordered"] is None
    assert after_b["field_versions"]["date_ordered"] == vb + 1  # a stale panel now sees a conflict


def test_a_po_with_invoices_cannot_be_cancelled(ws):
    c = ws["buyer"]
    a = _order(ws)
    id_ = _po(ws, a)["po_document_id"]
    c.post(f"/purchase-orders/{id_}/invoices", json={"order_id": a, "invoice_number": "I-1"})
    r = c.post(f"/purchase-orders/{id_}/cancel")
    assert r.status_code == 409 and _code(r) == "PO_HAS_INVOICES"
    assert c.get(f"/purchase-orders/{id_}").json()["status"] == "Open"


def test_a_cancelled_delivery_fee_order_is_replaced_not_stuck(ws):
    c = ws["buyer"]
    id_ = _po(ws, _order(ws))["po_document_id"]
    d1 = c.put(f"/purchase-orders/{id_}/delivery", json={"amount": "10.00"}).json()["delivery_order_id"]
    assert c.delete(f"/orders/{d1}").status_code in (200, 204)  # cancelled in the Orderbook
    po = c.get(f"/purchase-orders/{id_}").json()
    assert po["delivery_charge"] is None                        # the cancelled fee is no charge
    assert po["totals"]["total_ex"] == "100.00"
    r = c.put(f"/purchase-orders/{id_}/delivery", json={"amount": "20.00"})
    assert r.status_code == 200, r.text
    assert r.json()["delivery_order_id"] != d1 and r.json()["delivery_charge"] == "20.00"
    assert r.json()["totals"]["total_ex"] == "120.00"


def test_contact_emails_must_be_plain_addresses(ws):
    c, v = ws["buyer"], ws["v1"]
    for bad in ("a@b.com?bcc=evil@x.com", "a@b.com,evil@x.com", "a@b.com;x@y.com", "<a@b.com>", "a@b"):
        assert c.post(f"/suppliers/{v}/contacts", json={"name": "N", "email": bad}).status_code == 422, bad
    assert c.post(f"/suppliers/{v}/contacts", json={"name": "N", "email": "sam.o'neil+po@sub.example.com.au"}).status_code == 201


def test_excel_cells_from_order_text_are_never_formulas(ws):
    c = ws["buyer"]
    a = _order(ws, description='=HYPERLINK("http://evil","x")')
    id_ = _po(ws, a, memo="@SUM(1+1)")["po_document_id"]
    wb = load_workbook(io.BytesIO(c.get(f"/purchase-orders/{id_}/excel").content))
    bad = [(cell.coordinate, cell.value) for row in wb.active.iter_rows() for cell in row
           if cell.data_type == "f"]
    assert bad == []
    cells = [str(cell.value) for row in wb.active.iter_rows() for cell in row]
    assert '=HYPERLINK("http://evil","x")' in cells             # the text itself is kept
