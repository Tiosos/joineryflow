"""Starting the catalog from the orders export (E3, third slice): `app.importers.catalog_from_orders`.

Synthetic data only: the real export is customer data and never goes in the repo.
"""
import datetime as dt

from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import text

from app.db import SessionLocal
from app.importers import catalog_from_orders as co
from .helpers import login

pytestmark = pytest.mark.usefixtures("truncate_after")

D = dt.datetime


def _order(code, desc, supplier="Acme", when=D(2026, 3, 1), qty=1, total=10):
    return {"_PCode": code, "_BoardGroup": desc, "Supplier": supplier, "Date_Ordered": when, "QTY": qty,
            "_Cost_Sub": total}


def _xlsx(tmp_path, name, headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(headers))
    for r in rows:
        ws.append([r.get(h) for h in headers])
    path = tmp_path / name
    wb.save(path)
    return path


@pytest.mark.parametrize("code,desc,supplier,table", [
    ("708321", "3600 x 1800 x 18 - Raw Natural MR PB", "Boards Ltd", "board"),
    ("976411", "3600 x 1200 x 0.7 - ELEGANT OAK -[Finish: NATURAL]", None, "board"),
    ("114575", "100 x 21 x 1 - Elegant Oak UnGlued", None, "edging"),
    ("9088199", "Sensys angle hinge W90", "Hettich", "hardware"),
    ("RTR46ZM123000", "THREADED ROD 4.6 ZINC M12 X 3000MM", None, "hardware"),
    ("53236238", "1264 X 474 X 4 Silver Float FPAR", "Barton Glass", "glass"),
    ("34207667", "1264 X 474 X 4 Silver FPAR", None, "glass"),
    ("Tg_802932", "Solid Top - Tas Oak DAR 190 x 42 @ 2200mm", None, "timber"),                  # Solid Timber (2026-10-20)
    ("Tg_900001", "Solid Top - Tas Oak DAR 190 x 19 @ 2300mm", None, "timber"),
    ("ST-1", "Solid Timber Top 40mm Blackbutt", None, "timber"),
    ("ST-2", "Solid Oak slab 2400 x 600", None, "timber"),
    ("LT-1", "Solid Top - Laminate Oak DAR", None, "benchtop"),                                   # laminate / TG keep the old guess
    ("LT-2", "Solid Top - TG Solid Oak DAR", None, "benchtop"),
    ("SS-1", "Solid Top - Solid Surface White", None, "benchtop"),
    ("SS-2", "Solid Surface Corian Glacier White", None, "skip"),                                 # never timber, never guessed
    ("DD60D19", "Fisher Paykel Integrated Double DishDrawer Dishwasher", None, "appliance"),
    ("TG_DeliveryFee", "Delivery Fee", None, "skip"),
    ("Surcharge", "+0.50% Logistics Surcharge", None, "skip"),
    ("59877298", "Pcs:(3) 3870 x 3111.1 x 33 - ELEGANT OAK (Laminex) Square Edge", None, "skip"),   # made for one job
    ("46797947", "Toilet Partitions - As Per Attached Drawings", None, "skip"),
    ("103184", "Haulotte Star 6 Vertical Lift 01.02.2026 - 16.02.2026", None, "skip"),             # hire needs a project
    (None, "Delivery Fee", None, "skip"),
    ("ZZ1", "Walls Decorative Solid Panels", None, "skip"),                                       # could not tell
])
def test_the_guess_for_each_kind_of_product(code, desc, supplier, table):
    assert co.classify(code, desc, supplier)[0] == table


def test_propose_is_one_row_per_code_and_the_newest_order_names_it(tmp_path):
    rows = [_order("A1", "3600 x 1800 x 18 - Old wording MR PB", when=D(2026, 1, 1), qty=2, total=100),
            _order("A1", "3600 x 1800 x 18 - New wording MR PB", "Other Co", when=D(2026, 5, 1), qty=4, total=210),
            _order("G1", "100 x 100 x 4 Silver Float"), _order("G2", "100 x 100 x 4 Silver Float"),
            _order(None, "Delivery Fee")]
    products = {p.code: p for p in co.propose(rows)}
    assert len(products) == 4 and products["A1"].orders == 2
    assert products["A1"].description.endswith("New wording MR PB") and products["A1"].supplier == "Other Co"
    assert "other wording" in products["A1"].note and "more than one supplier" in products["A1"].note
    assert products["A1"].unit_cost == Decimal("52.5000")           # the latest order's total / qty, ex GST
    assert "same wording under 1 other code" in products["G1"].note
    assert products[None].table == "skip"


def test_a_latest_cost_of_zero_is_a_price_and_an_unusable_one_is_none():
    zero = co.propose([_order("Z", "Hinge", when=D(2026, 1, 1), total=30), _order("Z", "Hinge", when=D(2026, 2, 1), total=0)])
    assert zero[0].unit_cost == Decimal("0.0000")                   # settled: store 0, do not look further back
    no_qty = co.propose([_order("Q", "Hinge", qty=0, total=5), _order("R", "Hinge", qty=None, total=5)])
    assert [p.unit_cost for p in no_qty] == [None, None]


def test_the_review_sheet_round_trips_and_a_bad_table_is_refused(tmp_path):
    products = co.propose([_order("708321", "3600 x 1800 x 18 - Raw MR PB", "Boards Ltd"), _order("X", "Delivery Fee")])
    out = tmp_path / "review.xlsx"
    co.write_review(products, out)
    rows = co.read_review(out)
    assert [(r.code, r.table, r.supplier, r.unit_cost) for r in rows] == [
        ("708321", "board", "Boards Ltd", Decimal("10.0000")), ("X", "skip", "Acme", Decimal("10.0000"))]
    bad = _xlsx(tmp_path, "bad.xlsx", co.REVIEW_HEADERS,
                [{"Code": "A", "Description": "d", "Table": "paint"}])
    with pytest.raises(co.ImportErrorReport, match="paint"):
        co.read_review(bad)
    long_type = _xlsx(tmp_path, "long.xlsx", co.REVIEW_HEADERS,
                      [{"Code": "A", "Description": "d", "Table": "board", "Type": "x" * 65}])
    with pytest.raises(co.ImportErrorReport, match="type"):
        co.read_review(long_type)
    for cost in (-1, "abc"):
        worse = _xlsx(tmp_path, "worse.xlsx", co.REVIEW_HEADERS,
                      [{"Code": "A", "Description": "d", "Table": "board", "Unit cost": cost}])
        with pytest.raises(co.ImportErrorReport, match="unit cost"):
            co.read_review(worse)
    with pytest.raises(co.ImportErrorReport, match="missing columns"):
        co.read_orders(_xlsx(tmp_path, "o.xlsx", ["_PCode"], []))


def _review(tmp_path, rows):
    return co.read_review(_xlsx(tmp_path, "r.xlsx", co.REVIEW_HEADERS, rows))


def _apply(wid, uid, rows, *, commit=True, hire_project_id=None):
    db = SessionLocal()
    anomalies: list = []
    try:
        res = co.apply(db, rows, workspace_id=wid, actor_id=uid, source_sha256="z" * 64, anomalies=anomalies,
                       hire_project_id=hire_project_id)
        (db.commit if commit else db.rollback)()
        return res, anomalies
    finally:
        db.close()


def _q(sql, **p):
    db = SessionLocal()
    try:
        return db.execute(text(sql), p).all()
    finally:
        db.close()


def test_the_import_writes_each_table_links_suppliers_and_is_safe_to_rerun(tmp_path):
    _, wid, uid = login("admin")
    db = SessionLocal()
    try:
        vid = db.execute(text("INSERT INTO vendors(workspace_id, name, category) VALUES (:w, 'Boards Ltd', 'Other')"
                              " RETURNING vendor_id"), {"w": wid}).scalar()
        db.commit()
    finally:
        db.close()
    rows = _review(tmp_path, [
        {"Code": "B-1", "Description": "3600 x 1800 x 18 MR PB", "Supplier": "boards ltd", "Table": "board",
         "Unit cost": 45.5},
        {"Code": "H-1", "Description": "Hinge", "Supplier": "Nobody Ltd", "Table": "hardware", "Type": "Supplies"},
        {"Code": "C-1", "Description": "Cut glass", "Table": "custom_made"},
        {"Code": "T-1", "Description": "Solid top", "Table": "benchtop"},
        {"Code": "A-1", "Description": "Dishwasher", "Table": "appliance"},
        {"Code": "S-1", "Description": "Delivery", "Table": "skip"},
        {"Code": "B-1", "Description": "again", "Table": "board"},          # twice on the sheet
        {"Code": "L-" + "x" * 40, "Description": "too long for a board code", "Table": "board"},
    ])
    res, anomalies = _apply(wid, uid, rows, commit=False)
    assert dict(res.created) == {"board": 1, "hardware": 1, "custom_made": 1, "benchtop": 1, "appliance": 1}
    assert _q("SELECT 1 FROM board_materials WHERE sku = 'B-1'") == []              # a dry run writes nothing
    res, anomalies = _apply(wid, uid, rows)
    assert (res.skipped, res.linked, res.unlinked) == (1, 1, 1)
    assert {k for k, _ in anomalies} == {"duplicate_code", "code_too_long"}
    assert _q("SELECT code, sku, default_supplier, default_supplier_id, unit_cost FROM board_materials"
              " WHERE workspace_id = :w", w=wid) == [("B-1", "B-1", "boards ltd", vid, Decimal("45.5000"))]
    assert _q("SELECT unit_cost, product_type FROM hardware_materials WHERE workspace_id = :w", w=wid) == [
        (None, "Supplies")]                                          # no cost on the sheet; the old type is kept
    assert _q("SELECT product_type FROM board_materials WHERE workspace_id = :w", w=wid) == [(None,)]   # Type is optional
    assert _q("SELECT internal_ref, sku FROM custom_made WHERE workspace_id = :w", w=wid) == [("C-1", "C-1")]
    assert _q("SELECT slab_id FROM benchtop_materials WHERE workspace_id = :w", w=wid) == [("T-1",)]
    assert _q("SELECT model_number FROM appliances WHERE workspace_id = :w", w=wid) == [("A-1",)]
    assert _q("SELECT count(*) FROM audit_log WHERE event = 'import.catalog_from_orders'")[0][0] == 1

    again, _ = _apply(wid, uid, rows)
    assert dict(again.created) == {} and again.existing == 5


def test_a_code_another_workspace_already_uses_is_reported_not_crashed_on(tmp_path):
    _, wid, uid = login("admin")
    db = SessionLocal()
    try:
        other = db.execute(text("INSERT INTO workspace(slug, name) VALUES ('other-ws', 'Other') RETURNING id")).scalar()
        db.execute(text("INSERT INTO hardware_materials(workspace_id, sku, description) VALUES (:w, 'SHARED-1', 'theirs')"),
                   {"w": other})
        db.commit()
    finally:
        db.close()
    res, anomalies = _apply(wid, uid, _review(tmp_path, [{"Code": "SHARED-1", "Description": "mine", "Table": "hardware"}]))
    assert dict(res.created) == {} and [k for k, _ in anomalies] == ["code_used_elsewhere"]


def test_a_sheet_with_no_table_column_takes_the_table_from_the_type_and_hire_needs_a_project(tmp_path):
    _, wid, uid = login("admin")
    headers = ("Code", "Description", "Supplier", "Unit cost", "Type")        # column G (Table) deleted
    rows = co.read_review(_xlsx(tmp_path, "t.xlsx", headers, [
        {"Code": "G-1", "Description": "Cut glass", "Type": "Glass"},
        {"Code": "L-1", "Description": "Benchtop sheet", "Type": "laminate benchtop"},
        {"Code": "E-1", "Description": "Edge tape", "Type": "Edging"},
        {"Code": "T-1", "Description": "Oak board", "Type": "Solid Timber"},
        {"Code": "Q-1", "Description": "Scissor lift", "Type": "Equipment Hire"},
        {"Code": "N-1", "Description": "No type given"},
        {"Code": "U-1", "Description": "A type the list lacks", "Type": "Whatever"},
    ]))
    assert [r.table for r in rows] == ["glass", "benchtop", "edging", "timber", "hire", "skip", "skip"]
    res, anomalies = _apply(wid, uid, rows)                                # hire without a project is reported
    assert dict(res.created) == {"glass": 1, "benchtop": 1, "edging": 1, "timber": 1} and res.skipped == 2
    assert [k for k, _ in anomalies] == ["hire_needs_project"]
    assert _q("SELECT product_type FROM glass_materials WHERE workspace_id = :w", w=wid) == [("Glass",)]
    assert _q("SELECT sku, product_type FROM timber_materials WHERE workspace_id = :w", w=wid) == [("T-1", "Solid Timber")]
    db = SessionLocal()
    try:
        pid = db.execute(text("INSERT INTO projects(workspace_id, project_code, name) VALUES (:w, 'P-1', 'P')"
                              " RETURNING project_id"), {"w": wid}).scalar()
        db.commit()
    finally:
        db.close()
    res, anomalies = _apply(wid, uid, rows, hire_project_id=pid)
    assert res.created["hire"] == 1 and anomalies == []
    assert _q("SELECT contract_ref, project_id, product_type FROM equipment_hire WHERE workspace_id = :w", w=wid) == [
        ("Q-1", pid, "Equipment Hire")]


def test_type_decides_the_table_and_table_is_read_only_when_type_is_empty_or_unknown(tmp_path):
    rows = _review(tmp_path, [
        {"Code": "A", "Description": "d", "Type": "Glass", "Table": "hardware"},        # Type wins
        {"Code": "B", "Description": "d", "Type": "Edging", "Table": "skip"},          # Type wins over skip
        {"Code": "C", "Description": "d", "Type": "Supplies", "Table": "board"},       # known type, other table
        {"Code": "D", "Description": "d", "Type": "Whatever", "Table": "timber"},      # unknown Type: Table used
        {"Code": "E", "Description": "d", "Table": "benchtop"},                        # no Type: Table used
        {"Code": "F", "Description": "d", "Type": "Glass"},                            # no Table: Type
        {"Code": "G", "Description": "d", "Type": "Whatever"},                         # neither: skip
    ])
    assert [r.table for r in rows] == ["glass", "edging", "hardware", "timber", "benchtop", "glass", "skip"]


def test_the_read_me_sheet_says_type_wins(tmp_path):
    out = tmp_path / "review.xlsx"
    co.write_review(co.propose([_order("A", "3600 x 1800 x 18 - Raw MR PB")]), out)
    guide = " ".join(str(c.value) for row in openpyxl.load_workbook(out)["Read me"].iter_rows() for c in row)
    assert "it decides the table and Table is ignored" in guide and "only when Type is empty" in guide


def test_a_unit_cost_over_ten_whole_digits_is_refused_by_the_reader(tmp_path):
    ok = _review(tmp_path, [{"Code": "A", "Description": "d", "Table": "board", "Unit cost": "9999999999.9999"}])
    assert ok[0].unit_cost == Decimal("9999999999.9999")
    with pytest.raises(co.ImportErrorReport, match="unit cost"):
        _review(tmp_path, [{"Code": "A", "Description": "d", "Table": "board", "Unit cost": 10 ** 10}])


def _second_workspace(slug="other-ws"):
    db = SessionLocal()
    try:
        other = db.execute(text("INSERT INTO workspace(slug, name) VALUES (:s, 'Other') RETURNING id"),
                           {"s": slug}).scalar()
        db.commit()
        return other
    finally:
        db.close()


@pytest.mark.parametrize("type_,table,code_col,per_workspace", [
    ("edging", "edging_materials", "sku", True),
    ("glass", "glass_materials", "sku", True),
    ("timber", "timber_materials", "sku", True),
    ("hardware", "hardware_materials", "sku", False),
    ("board", "board_materials", "code", False),
    ("custom_made", "custom_made", "internal_ref", False),
])
def test_another_workspaces_code_blocks_only_where_the_database_constraint_is_global(
        tmp_path, type_, table, code_col, per_workspace):
    _, wid, uid = login("admin")
    other = _second_workspace()
    cols = {"board": "code, ", "custom_made": "internal_ref, "}.get(type_, "")
    vals = {"board": ":c, ", "custom_made": ":c, "}.get(type_, "")
    db = SessionLocal()
    try:
        db.execute(text(f"INSERT INTO {table}({cols}workspace_id, sku, description) VALUES ({vals}:w, :c, 'theirs')"),
                   {"w": other, "c": "SHARED-9"})
        db.commit()
    finally:
        db.close()
    rows = _review(tmp_path, [{"Code": "SHARED-9", "Description": "mine", "Table": type_}])
    res, anomalies = _apply(wid, uid, rows)
    if per_workspace:                    # UNIQUE (workspace_id, sku): the other workspace's SKU is no conflict
        assert res.created[type_] == 1 and anomalies == []
        assert _q(f"SELECT count(*) FROM {table} WHERE sku = 'SHARED-9'")[0][0] == 2
    else:                                # a global unique code: reported, not crashed on
        assert dict(res.created) == {} and [k for k, _ in anomalies] == ["code_used_elsewhere"]


def test_own_row_counts_as_existing_even_when_another_workspace_holds_the_same_code(tmp_path):
    _, wid, uid = login("admin")
    other = _second_workspace()
    db = SessionLocal()
    try:
        for w in (other, wid, other):       # the other workspace's row is first by id and second again
            if w == wid:
                db.execute(text("INSERT INTO glass_materials(workspace_id, sku, description) VALUES (:w, 'DUP', 'mine')"),
                           {"w": w})
            else:
                db.execute(text("INSERT INTO glass_materials(workspace_id, sku, description) VALUES (:w, 'DUP', 'x')"
                                " ON CONFLICT DO NOTHING"), {"w": w})
        db.commit()
    finally:
        db.close()
    res, anomalies = _apply(wid, uid, _review(tmp_path, [{"Code": "DUP", "Description": "d", "Table": "glass"}]))
    assert (res.existing, dict(res.created), anomalies) == (1, {}, [])
