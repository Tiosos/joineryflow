"""The PO's documents: the printable PO, the authorisation form and the Excel export.

`po_context` formats the money and dates once (dd/mm/yyyy like the old screen) so the templates stay
plain. The authorisation form is a printable sheet signed by hand: the software runs no approval
process (`0056`), so nothing here records a sign-off.
"""
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from ..printing.engine import render_template_to_pdf


def _money(v: Decimal | None) -> str:
    return "" if v is None else f"{v:,.2f}"


def _qty(v: Decimal | None) -> str:
    if v is None:
        return ""
    return f"{v.normalize():f}" if v == v.to_integral() else f"{v:f}".rstrip("0")


def _date(v: date | None) -> str:
    return "" if v is None else v.strftime("%d/%m/%Y")


def po_context(doc: dict) -> dict:
    lines = [
        {**line, "quantity_fmt": _qty(line["quantity"]), "unit_cost_fmt": _money(line["unit_cost"]),
         "grand_total_fmt": _money(line["grand_total"])}
        for line in doc["lines"]
    ]
    t = doc["totals"]
    return {"po": {
        **doc, "lines": lines,
        "date_ordered_fmt": _date(doc["date_ordered"]), "date_required_fmt": _date(doc["date_required"]),
        "date_due_fmt": _date(doc["date_due"]),
        "totals": {"total_ex_fmt": _money(t["total_ex"]), "gst_fmt": _money(t["gst"]),
                   "total_inc_fmt": _money(t["total_inc"])},
    }}


def po_pdf(doc: dict) -> bytes:
    return render_template_to_pdf("po.html", po_context(doc))


def authorisation_pdf(doc: dict) -> bytes:
    return render_template_to_pdf("po_authorisation.html", po_context(doc))


def po_xlsx(doc: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"PO {doc['po_no']}"
    bold = Font(bold=True)
    head = [
        ("PO number", doc["po_no"]), ("Supplier", doc["vendor_name"]),
        ("Reference no.", doc["reference_no"] or ""), ("Ordered by", doc["ordered_by_name"] or ""),
        ("Date ordered", _date(doc["date_ordered"])), ("Date required", _date(doc["date_required"])),
        ("Date due", _date(doc["date_due"])), ("Deliver to", doc["deliver_to"] or ""),
        ("Memo to supplier", doc["memo"] or ""),
    ]
    for label, value in head:
        ws.append([label, value])
        ws.cell(row=ws.max_row, column=1).font = bold
    ws.append([])
    columns = ["#", "Qty", "Item", "Pcode", "Location", "Project", "Cost", "Um", "Amount (ex GST)",
               "GST", "Amount (inc GST)", "Status", "Invoiced"]
    ws.append(columns)
    for c in range(1, len(columns) + 1):
        ws.cell(row=ws.max_row, column=c).font = bold
    for l in doc["lines"]:
        ws.append([
            l["line_no"], _num(l["quantity"]), l["description"], l["product_code"] or "",
            l["location"] or "", l["project_name"] or "", _num(l["unit_cost"]),
            l["unit_of_measure"] or "", _num(l["total_amount"]), _num(l["gst_amount"]),
            _num(l["grand_total"]), l["status"], "yes" if l["invoiced"] else "",
        ])
    ws.append([])
    t = doc["totals"]
    for label, value in (("Total Order", t["total_ex"]), ("Gst", t["gst"]), ("Gst inc", t["total_inc"])):
        ws.append(["", "", "", "", "", "", "", "", label, "", _num(value)])
        ws.cell(row=ws.max_row, column=9).font = bold
    for col, width in zip("ABCDEFGHIJKLM", (6, 8, 48, 14, 18, 22, 10, 6, 16, 10, 16, 12, 9)):
        ws.column_dimensions[col].width = width
    ws["B2"].alignment = Alignment(horizontal="left")
    # order text is typed by people: a value starting with = + - @ must stay text, never a formula
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _num(v: Decimal | None) -> float | None:
    return None if v is None else float(v)
